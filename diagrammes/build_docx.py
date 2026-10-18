# -*- coding: utf-8 -*-
"""Dessine les diagrammes UML et les insère dans un document Word."""
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from pathlib import Path

OUT = Path(__file__).resolve().parent
FONT = "C:/Windows/Fonts/calibri.ttf"
FONTB = "C:/Windows/Fonts/calibrib.ttf"
INK = "#1c2430"
MUTED = "#5c6b7a"
LINE = "#243044"
FILL = "#ffffff"
ACTOR = "#1f4e79"
UC = "#eef4f8"
BAND = "#f7f9fb"


def font(size, bold=False):
    return ImageFont.truetype(FONTB if bold else FONT, size)


def text_size(draw, text, f):
    box = draw.textbbox((0, 0), text, font=f)
    return box[2] - box[0], box[3] - box[1]


def draw_ellipse(draw, box, fill, outline):
    draw.ellipse(box, fill=fill, outline=outline, width=2)


def actor(draw, cx, cy, label):
    f = font(16, True)
    r = 16
    draw.ellipse((cx - r, cy - 52, cx + r, cy - 20), outline=ACTOR, width=2)
    draw.line((cx, cy - 20, cx, cy + 18), fill=ACTOR, width=2)
    draw.line((cx - 22, cy - 4, cx + 22, cy - 4), fill=ACTOR, width=2)
    draw.line((cx, cy + 18, cx - 18, cy + 48), fill=ACTOR, width=2)
    draw.line((cx, cy + 18, cx + 18, cy + 48), fill=ACTOR, width=2)
    tw, _ = text_size(draw, label, f)
    draw.text((cx - tw / 2, cy + 54), label, fill=INK, font=f)


def usecase(draw, cx, cy, label, w=210, h=46):
    f = font(15)
    box = (cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2)
    draw.rounded_rectangle(box, radius=22, fill=UC, outline=LINE, width=2)
    tw, th = text_size(draw, label, f)
    draw.text((cx - tw / 2, cy - th / 2 - 1), label, fill=INK, font=f)
    return (cx, cy)


def edge(draw, a, b):
    draw.line((a[0], a[1], b[0], b[1]), fill="#8aa0b4", width=2)


def build_usecase():
    img = Image.new("RGB", (1680, 1100), "white")
    draw = ImageDraw.Draw(img)
    title = font(28, True)
    draw.text((40, 24), "Diagramme de cas d'utilisation", fill=INK, font=title)
    sub = font(16)
    draw.text((40, 64), "Gestion des stages  ·  acteurs et cas réellement couverts par l'application", fill=MUTED, font=sub)

    actors = {
        "etu": (130, 250, "Étudiant"),
        "ent": (130, 500, "Entreprise"),
        "adm": (130, 730, "Super admin"),
        "chef": (130, 860, "Chef de département"),
    }
    for _, y, label in actors.values():
        actor(draw, 130, y, label)

    login = usecase(draw, 430, 160, "Se connecter", 180, 44)
    for key in actors:
        edge(draw, (170, actors[key][1] - 10), (login[0] - 90, login[1]))

    etu_cases = [
        (700, 140, "Consulter les offres"),
        (960, 140, "Gérer le profil"),
        (1220, 140, "Gérer le CV"),
        (1480, 140, "Passer le QCM"),
        (700, 220, "Postuler"),
        (960, 220, "Voir ses candidatures"),
        (1220, 220, "Voir les recommandations"),
        (1480, 220, "Lire les notifications"),
        (700, 300, "Tenir le journal"),
        (960, 300, "Réaliser les tâches"),
        (1220, 300, "Ouvrir une réclamation"),
        (1480, 300, "Voir le calendrier"),
        (960, 380, "Signer la clôture"),
    ]
    etu_pos = []
    for x, y, label in etu_cases:
        etu_pos.append(usecase(draw, x, y, label))
        edge(draw, (168, actors["etu"][1]), (x - 105, y))

    ent_cases = [
        (700, 470, "Gérer les offres"),
        (960, 470, "Voir les candidatures"),
        (1220, 470, "Accepter ou refuser"),
        (1480, 470, "Assigner une tâche"),
        (700, 550, "Valider ou rejeter"),
        (960, 550, "Traiter les réclamations"),
        (1220, 550, "Signer la clôture"),
    ]
    for x, y, label in ent_cases:
        usecase(draw, x, y, label)
        edge(draw, (168, actors["ent"][1]), (x - 105, y))

    for x, y, label in (
        (700, 700, "Gérer les entreprises"),
        (980, 700, "Valider une entreprise"),
        (1260, 700, "Gérer les étudiants"),
    ):
        usecase(draw, x, y, label, 230)
        edge(draw, (168, actors["adm"][1]), (x - 115, y))

    chef_uc = usecase(draw, 780, 860, "Accéder à l'espace administration", 280, 48)
    edge(draw, (168, actors["chef"][1]), (chef_uc[0] - 140, chef_uc[1]))

    note = font(14)
    draw.text(
        (40, 1040),
        "Inclusions : postuler dépend de la consultation des offres ; accepter ou refuser dépend des candidatures ; les recommandations s'appuient sur le profil.",
        fill=MUTED,
        font=note,
    )
    path = OUT / "cas-utilisation.png"
    img.save(path, "PNG")
    return path


def class_box(draw, x, y, name, fields, w=250):
    fb = font(15, True)
    f = font(13)
    line_h = 18
    head = 30
    body = 10 + line_h * max(1, len(fields)) + 8
    h = head + body
    draw.rectangle((x, y, x + w, y + h), fill=FILL, outline=LINE, width=2)
    draw.rectangle((x, y, x + w, y + head), fill="#e7eef5", outline=LINE, width=2)
    tw, th = text_size(draw, name, fb)
    draw.text((x + (w - tw) / 2, y + (head - th) / 2), name, fill=INK, font=fb)
    yy = y + head + 8
    for field in fields:
        draw.text((x + 10, yy), field, fill=INK, font=f)
        yy += line_h
    return {"x": x, "y": y, "w": w, "h": h, "cx": x + w / 2, "cy": y + h / 2}


def connect(draw, a, b, label, side_a="right", side_b="left"):
    def pt(box, side):
        if side == "right":
            return box["x"] + box["w"], box["cy"]
        if side == "left":
            return box["x"], box["cy"]
        if side == "bottom":
            return box["cx"], box["y"] + box["h"]
        return box["cx"], box["y"]

    p1 = pt(a, side_a)
    p2 = pt(b, side_b)
    draw.line((p1, p2), fill="#5d7386", width=2)
    mx, my = (p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2
    f = font(12)
    if label:
        tw, th = text_size(draw, label, f)
        draw.rectangle((mx - tw / 2 - 3, my - th / 2 - 2, mx + tw / 2 + 3, my + th / 2 + 2), fill="white")
        draw.text((mx - tw / 2, my - th / 2), label, fill=MUTED, font=f)


def build_classes():
    img = Image.new("RGB", (1880, 1280), "white")
    draw = ImageDraw.Draw(img)
    draw.text((40, 20), "Diagramme de classes", fill=INK, font=font(28, True))
    draw.text(
        (40, 58),
        "Entités JPA du backend. Le stage suivi est une candidature acceptée.",
        fill=MUTED,
        font=font(16),
    )

    ent = class_box(draw, 40, 110, "Entreprise", [
        "+ nom : String",
        "+ emailContact : String",
        "+ taille : TailleEntreprise",
        "+ statutValidation",
        "+ motifRejet : String",
    ])
    user = class_box(draw, 430, 110, "Utilisateur", [
        "+ id : Long",
        "+ nom, prenom : String",
        "+ email : String",
        "+ role : Role",
        "+ actif : boolean",
        "+ filiere : String",
        "+ niveauEtudes : String",
        "+ competences : List",
    ], 270)
    offre = class_box(draw, 40, 360, "OffreStage", [
        "+ titre : String",
        "+ domaine : String",
        "+ mode : ModeTravail",
        "+ dateLimite : LocalDate",
        "+ nombrePlaces : Integer",
        "+ statut : StatutOffre",
        "+ competences : List",
    ])
    cand = class_box(draw, 430, 430, "Candidature", [
        "+ statut : StatutCandidature",
        "+ statutStage : StatutStage",
        "+ message : String",
        "+ signatureEtudiantPath",
        "+ signatureEncadrantPath",
        "+ dateCloture : LocalDateTime",
    ], 280)
    tache = class_box(draw, 40, 680, "Tache", [
        "+ date : LocalDate",
        "+ titre : String",
        "+ statutApprobation",
        "+ commentaireEncadrant",
    ])
    tass = class_box(draw, 430, 720, "TacheAssignee", [
        "+ titre : String",
        "+ dateEcheance : LocalDate",
        "+ statut : StatutTacheAssignee",
    ], 280)
    notif = class_box(draw, 820, 110, "NotificationEtudiant", [
        "+ type : TypeNotification",
        "+ message : String",
        "+ lue : boolean",
    ], 270)
    recl = class_box(draw, 820, 360, "Reclamation", [
        "+ typeReclamation",
        "+ objet : String",
        "+ statut : StatutReclamation",
    ], 270)
    msg = class_box(draw, 820, 600, "MessageReclamation", [
        "+ contenu : String",
        "+ dateEnvoi : LocalDateTime",
    ], 270)
    qcm = class_box(draw, 1220, 110, "Qcm", [
        "+ titre : String",
        "+ noteMinimalePassage",
        "+ actif : boolean",
    ], 250)
    tent = class_box(draw, 1560, 110, "TentativeQcm", [
        "+ scorePourcentage",
        "+ reussi : boolean",
        "+ numeroTentative",
    ], 270)
    question = class_box(draw, 1220, 360, "Question", [
        "+ enonce : String",
        "+ type : TypeQuestion",
        "+ points : Integer",
    ], 250)
    option = class_box(draw, 1560, 360, "OptionReponse", [
        "+ libelle : String",
        "+ estCorrecte : boolean",
    ], 250)

    connect(draw, ent, user, "1  —  0..*  compte", "right", "left")
    connect(draw, ent, offre, "1  —  1..*", "bottom", "top")
    connect(draw, offre, cand, "1  —  0..*", "right", "left")
    connect(draw, user, cand, "etudiant  1 — 0..*", "bottom", "top")
    connect(draw, cand, tache, "journal  1 — 0..*", "left", "right")
    connect(draw, cand, tass, "1  —  0..*", "bottom", "top")
    connect(draw, user, notif, "1  —  0..*", "right", "left")
    connect(draw, cand, recl, "1  —  0..*", "right", "left")
    connect(draw, recl, msg, "1  —  1..*", "bottom", "top")
    connect(draw, qcm, tent, "1  —  0..*", "right", "left")
    connect(draw, qcm, question, "1  —  1..*", "bottom", "top")
    connect(draw, question, option, "1  —  2..*", "right", "left")

    f = font(14)
    notes = [
        "Candidature est unique pour le couple offre + étudiant.",
        "Candidature.encadrant et Reclamation.traitePar pointent vers Utilisateur (0..1).",
        "TacheAssignee.encadrant et MessageReclamation.auteur pointent vers Utilisateur.",
        "NotificationEtudiant référence la candidature (obligatoire) et, au plus, une tâche assignée.",
        "TentativeQcm référence l'étudiant et le QCM. Le service Python de recommandation n'est pas persisté.",
    ]
    y = 1040
    for line in notes:
        draw.text((40, y), "•  " + line, fill=INK, font=f)
        y += 28
    path = OUT / "classes.png"
    img.save(path, "PNG")
    return path


def add_picture(doc, path, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=width)


def main():
    uc = build_usecase()
    cl = build_classes()
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)

    title = doc.add_heading("Gestion des stages", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    intro = doc.add_paragraph(
        "Diagrammes établis à partir des rôles, des routes Angular et des entités JPA du projet."
    )
    intro.runs[0].font.size = Pt(12)
    intro.runs[0].font.color.rgb = RGBColor(0x5C, 0x6B, 0x7A)

    doc.add_heading("1. Diagramme de cas d'utilisation", level=1)
    add_picture(doc, uc, Cm(26.5))
    doc.add_paragraph(
        "L'entreprise dont le compte n'est pas validé ne peut pas se connecter. "
        "Le chef de département accède à l'espace administration, sans créer d'entreprise ni d'étudiant. "
        "Ces deux actions sont réservées au super administrateur."
    )

    doc.add_page_break()
    doc.add_heading("2. Diagramme de classes", level=1)
    add_picture(doc, cl, Cm(26.5))

    target = OUT.parent / "diagrammes-gestion-stages.docx"
    doc.save(target)
    print(target)


if __name__ == "__main__":
    main()
